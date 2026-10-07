"""Word and PowerPoint uploads, and OCR for scanned PDFs.

Office files are built in the tests with python-docx / python-pptx, scanned PDFs
with PyMuPDF (a page rendered to a picture and placed in a new PDF, so it has an
image and no text layer). The OCR logic is tested with a stand-in engine; one test
runs the real engine and is skipped when it isn't installed.
"""

import io
import zipfile

import fitz
import pytest

from app import models
from app.services import ocr_service, text_extraction
from app.services.text_extraction import (
    NO_OCR_MESSAGE,
    UploadRejected,
    extract_docx_pages,
    extract_pages,
    extract_pdf_pages,
    extract_pptx_pages,
)
from helpers import make_topic


# --- builders ----------------------------------------------------------------------------------------


def docx_bytes(build=None):
    from docx import Document

    document = Document()
    if build:
        build(document)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def pptx_bytes(build):
    from pptx import Presentation

    presentation = Presentation()
    build(presentation)
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


def add_slide(presentation, title=None, body=None, notes=None, layout=1):
    slide = presentation.slides.add_slide(presentation.slide_layouts[layout])
    if title is not None:
        slide.shapes.title.text = title
    if body is not None and len(slide.placeholders) > 1:
        slide.placeholders[1].text = body
    if notes is not None:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def scanned_pdf(page_texts, text_layer_pages=()):
    """A PDF whose pages are pictures of `page_texts`; pages listed in
    `text_layer_pages` (0-based) keep a real text layer instead."""
    source = fitz.open()
    for text in page_texts:
        page = source.new_page()
        page.insert_text((72, 100), text, fontsize=20)
    out = fitz.open()
    for index, page in enumerate(source):
        new_page = out.new_page(width=page.rect.width, height=page.rect.height)
        if index in text_layer_pages:
            new_page.insert_text((72, 100), page_texts[index], fontsize=20)
        else:
            new_page.insert_image(new_page.rect, pixmap=page.get_pixmap(dpi=150))
    data = out.tobytes()
    out.close()
    source.close()
    return data


def upload(client, topic_id, filename, content):
    return client.post(
        "/sources/upload",
        data={"topic_id": str(topic_id), "source_type": "self_supplied"},
        files={"file": (filename, content, "application/octet-stream")},
    )


@pytest.fixture()
def topic(client):
    return make_topic(client, "Databases", course="CS")["id"]


# --- Word ----------------------------------------------------------------------------------------------


def test_a_docx_gives_its_paragraphs_and_headings_in_order_as_one_untagged_page():
    def build(doc):
        doc.add_heading("Normalization", level=1)
        doc.add_paragraph("Normalization removes redundant data from tables.")
        doc.add_heading("Normal forms", level=2)
        doc.add_paragraph("First normal form needs atomic values.")

    ((number, text),) = extract_docx_pages(docx_bytes(build))
    assert number is None
    assert text.index("Normalization\n") < text.index("removes redundant") < text.index("Normal forms") < text.index("atomic values")


def test_a_docx_table_comes_out_a_row_per_line_in_its_place():
    def build(doc):
        doc.add_paragraph("Before the table.")
        table = doc.add_table(rows=2, cols=2)
        for r, row in enumerate([["Method", "Cost"], ["Chaining", "O(1)"]]):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
        doc.add_paragraph("After the table.")

    ((_, text),) = extract_docx_pages(docx_bytes(build))
    assert "Method | Cost\nChaining | O(1)" in text
    assert text.index("Before") < text.index("Method | Cost") < text.index("After")


def test_a_merged_table_cell_is_not_repeated():
    def build(doc):
        table = doc.add_table(rows=1, cols=3)
        merged = table.cell(0, 0).merge(table.cell(0, 1))
        merged.text = "Wide header"
        table.cell(0, 2).text = "Other"

    ((_, text),) = extract_docx_pages(docx_bytes(build))
    assert text == "Wide header | Other"


def test_an_empty_docx_has_no_pages():
    assert extract_docx_pages(docx_bytes()) == []


def test_a_docx_upload_is_chunked_and_searchable(client, db_session, topic):
    def build(doc):
        doc.add_heading("Hash tables", level=1)
        doc.add_paragraph("A hash table maps keys to values using a hash function to pick a bucket.")

    resp = upload(client, topic, "notes.docx", docx_bytes(build))
    assert resp.status_code == 200, resp.text
    assert resp.json()["title"] == "notes.docx"
    chunks = db_session.query(models.Chunk).filter(models.Chunk.topic_id == topic).all()
    assert chunks and all(c.page_number is None for c in chunks)
    assert "hash function" in " ".join(c.chunk_text for c in chunks)


def test_a_docx_extension_on_something_else_is_415(client, db_session, topic):
    resp = upload(client, topic, "fake.docx", b"just some text pretending to be Word")
    assert resp.status_code == 415 and "isn't a Word document" in resp.json()["detail"]
    assert db_session.query(models.Source).count() == 0


def test_a_zip_that_is_not_a_docx_is_415(client, topic):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("readme.txt", "hello")
    resp = upload(client, topic, "archive.docx", buffer.getvalue())
    assert resp.status_code == 415 and "isn't a Word document" in resp.json()["detail"]


def test_a_pptx_is_not_accepted_as_a_docx(client, topic):
    data = pptx_bytes(lambda p: add_slide(p, "Title"))
    assert upload(client, topic, "slides.docx", data).status_code == 415


def test_a_corrupt_docx_is_422(client, topic):
    resp = upload(client, topic, "broken.docx", b"PK\x03\x04" + b"this is not a real archive" * 20)
    assert resp.status_code == 422 and "corrupted" in resp.json()["detail"]


def test_a_docx_that_expands_too_much_is_refused_before_it_is_parsed(client, topic, monkeypatch):
    monkeypatch.setattr(text_extraction, "MAX_UNCOMPRESSED_BYTES", 1000)
    parsed = []
    monkeypatch.setattr(text_extraction, "extract_docx_pages", lambda raw: parsed.append(1) or [])
    resp = upload(client, topic, "bomb.docx", docx_bytes(lambda d: d.add_paragraph("x" * 50)))
    assert resp.status_code == 413 and "expands to too much" in resp.json()["detail"]
    assert parsed == []


def test_a_docx_with_no_text_is_422(client, topic):
    resp = upload(client, topic, "empty.docx", docx_bytes())
    assert resp.status_code == 422 and "No text could be extracted" in resp.json()["detail"]


def test_old_binary_office_formats_are_still_415(client, topic):
    for name in ("old.doc", "old.ppt", "old.xls"):
        assert upload(client, topic, name, b"\xd0\xcf\x11\xe0 legacy").status_code == 415


# --- PowerPoint -----------------------------------------------------------------------------------------


def test_each_slide_is_a_page_numbered_by_its_slide():
    def build(prs):
        add_slide(prs, "Hash tables", "Keys map to buckets.", notes="Mention the load factor.")
        add_slide(prs, "Sorting", "Merge sort halves the list.")

    pages = extract_pptx_pages(pptx_bytes(build))
    assert [number for number, _ in pages] == [1, 2]
    first, second = pages[0][1], pages[1][1]
    assert first.startswith("Hash tables") and "Keys map to buckets." in first
    assert first.endswith("Speaker notes: Mention the load factor.")
    assert second.startswith("Sorting") and "Speaker notes" not in second


def test_a_slide_with_no_text_is_skipped_but_numbering_is_kept():
    def build(prs):
        add_slide(prs, "First", "one")
        prs.slides.add_slide(prs.slide_layouts[6])  # blank layout, nothing on it
        add_slide(prs, "Third", "three")

    assert [number for number, _ in extract_pptx_pages(pptx_bytes(build))] == [1, 3]


def test_a_table_and_a_grouped_text_box_on_a_slide_are_read():
    from pptx.util import Inches

    def build(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
        slide.shapes.title.text = "Costs"
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
        for r, row in enumerate([["Op", "Time"], ["Lookup", "O(1)"]]):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
        group = slide.shapes.add_group_shape()
        box = group.shapes.add_textbox(Inches(1), Inches(4), Inches(3), Inches(1))
        box.text_frame.text = "Inside a group"

    ((number, text),) = extract_pptx_pages(pptx_bytes(build))
    assert number == 1
    assert "Op | Time\nLookup | O(1)" in text and "Inside a group" in text


def test_a_pptx_upload_keeps_slide_numbers_as_page_numbers(client, db_session, topic):
    def build(prs):
        add_slide(prs, "Hash tables", "A hash function maps a key to a bucket index.")
        add_slide(prs, "Sorting", "Merge sort splits the list and merges sorted halves.")

    resp = upload(client, topic, "lecture.pptx", pptx_bytes(build))
    assert resp.status_code == 200, resp.text
    chunks = db_session.query(models.Chunk).filter(models.Chunk.topic_id == topic).order_by(models.Chunk.chunk_index).all()
    assert [c.page_number for c in chunks] == [1, 2]
    assert "Merge sort" in chunks[1].chunk_text and "Merge sort" not in chunks[0].chunk_text


def test_a_pptx_extension_on_something_else_is_415(client, topic):
    assert upload(client, topic, "fake.pptx", b"plain text").status_code == 415
    assert upload(client, topic, "doc.pptx", docx_bytes(lambda d: d.add_paragraph("hi"))).status_code == 415


def test_a_pptx_with_no_text_is_422(client, topic):
    data = pptx_bytes(lambda prs: prs.slides.add_slide(prs.slide_layouts[6]))
    assert upload(client, topic, "empty.pptx", data).status_code == 422


# --- OCR: which pages, how many, and when it is unavailable ---------------------------------------------


class FakeOCR:
    def __init__(self, text="Recognised words from the scan."):
        self.text = text
        self.calls = 0

    def __call__(self, image):
        self.calls += 1
        assert image.ndim == 3 and image.shape[2] == 3  # a height x width x RGB picture
        return self.text


def test_a_scanned_page_is_read_with_ocr_and_keeps_its_page_number():
    ocr = FakeOCR()
    pages = extract_pdf_pages(scanned_pdf(["one", "two", "three"]), recognize=ocr, ocr_enabled=True)
    assert [n for n, _ in pages] == [1, 2, 3]
    assert all(text == "Recognised words from the scan." for _, text in pages)
    assert ocr.calls == 3


def test_pages_with_a_text_layer_are_not_sent_to_ocr():
    ocr = FakeOCR()
    data = scanned_pdf(["A real page with plenty of text on it, enough to count.", "scanned page"], text_layer_pages=(0,))
    pages = extract_pdf_pages(data, recognize=ocr, ocr_enabled=True)
    assert ocr.calls == 1
    assert pages[0] == (1, "A real page with plenty of text on it, enough to count.")
    assert pages[1] == (2, "Recognised words from the scan.")


def test_a_short_text_layer_on_a_picture_page_keeps_both():
    ocr = FakeOCR("Text read from the picture.")
    source = fitz.open()
    source.new_page().insert_text((72, 100), "scanned body", fontsize=20)
    out = fitz.open()
    page = out.new_page()
    page.insert_image(page.rect, pixmap=source[0].get_pixmap(dpi=150))
    page.insert_text((72, 40), "12", fontsize=10)  # only a page number is real text
    data = out.tobytes()
    ((number, text),) = extract_pdf_pages(data, recognize=ocr, ocr_enabled=True)
    assert number == 1 and text.startswith("12") and text.endswith("Text read from the picture.")


def test_ocr_that_finds_nothing_drops_the_page_instead_of_adding_blank_text():
    pages = extract_pdf_pages(scanned_pdf(["one", "two"]), recognize=FakeOCR(""), ocr_enabled=True)
    assert pages == []


def test_a_pdf_with_text_and_no_pictures_never_touches_ocr(sample_pdf_bytes):
    ocr = FakeOCR()
    pages = extract_pdf_pages(sample_pdf_bytes, recognize=ocr, ocr_enabled=True)
    assert [n for n, _ in pages] == [1, 2] and ocr.calls == 0


def test_a_fully_scanned_pdf_without_ocr_says_how_to_enable_it():
    with pytest.raises(UploadRejected) as info:
        extract_pdf_pages(scanned_pdf(["one"]), ocr_enabled=False)
    assert info.value.status_code == 422 and info.value.detail == NO_OCR_MESSAGE
    assert "requirements-ocr.txt" in NO_OCR_MESSAGE


def test_a_mixed_pdf_without_ocr_keeps_its_text_pages():
    data = scanned_pdf(["A real page with plenty of text on it, enough to count.", "scanned"], text_layer_pages=(0,))
    pages = extract_pdf_pages(data, ocr_enabled=False)
    assert [n for n, _ in pages] == [1]


def test_too_many_scanned_pages_is_a_413_that_names_the_setting(monkeypatch):
    monkeypatch.setattr(text_extraction, "OCR_MAX_PAGES", 2)
    ocr = FakeOCR()
    with pytest.raises(UploadRejected) as info:
        extract_pdf_pages(scanned_pdf(["a", "b", "c"]), recognize=ocr, ocr_enabled=True)
    assert info.value.status_code == 413 and "OCR_MAX_PAGES" in info.value.detail
    assert ocr.calls == 0  # refused before spending minutes on it


def test_the_upload_route_reads_a_scanned_pdf_through_ocr(client, db_session, topic, monkeypatch):
    monkeypatch.setattr(ocr_service, "ocr_available", lambda: True)
    monkeypatch.setattr(ocr_service, "recognize_page", lambda image: "Scanned lecture about hash tables and buckets.")
    resp = upload(client, topic, "scan.pdf", scanned_pdf(["a", "b"]))
    assert resp.status_code == 200, resp.text
    chunks = db_session.query(models.Chunk).filter(models.Chunk.topic_id == topic).order_by(models.Chunk.chunk_index).all()
    assert [c.page_number for c in chunks] == [1, 2]
    assert "hash tables" in chunks[0].chunk_text


def test_the_upload_route_explains_a_scan_when_ocr_is_missing(client, db_session, topic, monkeypatch):
    monkeypatch.setattr(ocr_service, "ocr_available", lambda: False)
    resp = upload(client, topic, "scan.pdf", scanned_pdf(["a"]))
    assert resp.status_code == 422 and "scanned PDF" in resp.json()["detail"]
    assert db_session.query(models.Source).count() == 0


# --- the OCR service itself ----------------------------------------------------------------------------


def box(left, top, width=100, height=20):
    return [[left, top], [left + width, top], [left + width, top + height], [left, top + height]]


def test_ocr_lines_come_out_top_to_bottom_then_left_to_right():
    results = [
        (box(300, 12), "right", 0.9),
        (box(10, 60), "second row", 0.9),
        (box(10, 10), "left", 0.9),
    ]
    assert ocr_service._reading_order(results) == ["left", "right", "second row"]


def test_ocr_lines_the_engine_is_unsure_of_are_dropped():
    results = [(box(0, 0), "clear", 0.95), (box(0, 40), "smudge", 0.2), (box(0, 80), "", 0.99)]
    assert ocr_service._reading_order(results) == ["clear"]


def test_ocr_can_be_switched_off(monkeypatch):
    monkeypatch.setattr(ocr_service, "OCR_ENABLED", False)
    assert ocr_service.ocr_available() is False


def test_ocr_is_unavailable_when_the_engine_is_not_installed(monkeypatch):
    import builtins

    real_import = builtins.__import__

    def refuse(name, *args, **kwargs):
        if name.startswith("rapidocr_onnxruntime"):
            raise ImportError("not installed")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(ocr_service, "_import_failed", False)
    monkeypatch.setattr(builtins, "__import__", refuse)
    assert ocr_service.ocr_available() is False
    monkeypatch.setattr(builtins, "__import__", real_import)
    monkeypatch.setattr(ocr_service, "_import_failed", False)


def test_the_real_engine_reads_a_scanned_page():
    pytest.importorskip("rapidocr_onnxruntime")
    data = scanned_pdf(["Database normalization removes redundant data.", "Third normal form has no transitive dependencies."])
    pages = extract_pdf_pages(data, ocr_enabled=True)
    text = " ".join(t for _, t in pages).lower()
    assert [n for n, _ in pages] == [1, 2]
    assert "normalization" in text and "redundant" in text
    assert "transitive" in text and "dependencies" in text


def test_extract_pages_dispatches_office_and_scanned_files_by_extension():
    assert extract_pages("a.docx", docx_bytes(lambda d: d.add_paragraph("Word text")))[0][1] == "Word text"
    slides = extract_pages("a.pptx", pptx_bytes(lambda p: add_slide(p, "Slide title", "Slide text")))
    assert slides[0][0] == 1
    with pytest.raises(UploadRejected):
        extract_pages("a.rtf", b"{\\rtf1}")
